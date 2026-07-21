from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Dict, Iterable, List


def normalize_key(value: str) -> str:
    value = str(value or "").strip().lower()
    value = value.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    value = re.sub(r"[^a-z0-9]+", "", value)
    return value


def field_name_map(fields: Iterable[dict]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for field in fields:
        name = field.get("name", "")
        if not name:
            continue
        result[normalize_key(name)] = name

        # Helpful variants for common VADAFOK field naming.
        result[normalize_key(name.replace("_", " "))] = name
        result[normalize_key(name.replace("-", " "))] = name
    return result


def analyze_columns(rows: List[Dict[str, str]], fields: List[dict]) -> tuple[list[str], list[str]]:
    """Return input columns that match template fields and those ignored."""
    if not rows:
        return [], []
    fmap = field_name_map(fields)
    output_columns = {"outputname", "filename", "file"}
    columns = list(dict.fromkeys(str(column) for row in rows for column in row))
    matched = [column for column in columns if normalize_key(column) in fmap]
    ignored = [
        column for column in columns
        if normalize_key(column) not in fmap
        and normalize_key(column) not in output_columns
    ]
    return matched, ignored


def unique_output_name(base_name: str, used_names: Iterable[str]) -> str:
    """Return a readable case-insensitive unique name for a batch item."""
    base = str(base_name or "card").strip() or "card"
    used = {str(name).strip().casefold() for name in used_names}
    if base.casefold() not in used:
        return base
    number = 2
    while f"{base}_{number}".casefold() in used:
        number += 1
    return f"{base}_{number}"


def read_csv(path: Path) -> List[Dict[str, str]]:
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]
    last_error = None

    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;	")
                except Exception:
                    dialect = csv.excel
                reader = csv.DictReader(f, dialect=dialect)
                rows = [
                    {str(k or "").strip(): str(v or "").strip() for k, v in row.items()}
                    for row in reader
                ]
                return [row for row in rows if any(row.values())]
        except Exception as e:
            last_error = e

    raise last_error or RuntimeError("CSV konnte nicht gelesen werden.")


def read_xlsx(path: Path) -> List[Dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    try:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(v or "").strip() for v in rows[0]]
        result = []

        for row in rows[1:]:
            item = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else ""
                item[header] = "" if value is None else str(value).strip()
            if any(v for v in item.values()):
                result.append(item)

        return result
    finally:
        wb.close()


def read_table(path: Path) -> List[Dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError(f"Nicht unterstütztes Format: {suffix}")


def rows_to_batch_items(rows: List[Dict[str, str]], template_name: str, fields: List[dict], output_prefix: str, profile: str) -> List[dict]:
    fmap = field_name_map(fields)
    items = []

    for idx, row in enumerate(rows, start=1):
        values: Dict[str, str] = {}

        for column, value in row.items():
            normalized = normalize_key(column)
            if normalized in fmap:
                values[fmap[normalized]] = value

        # Output name can be supplied by several common column names.
        output_name = next(
            (
                value for column, value in row.items()
                if normalize_key(column) in {"outputname", "filename", "file"}
                and str(value).strip()
            ),
            f"{output_prefix}_{idx:03d}",
        )

        items.append({
            "template": template_name,
            "output_name": output_name,
            "profile": profile,
            "values": values,
        })

    return items



def batch_projects_dir() -> Path:
    folder = Path.cwd() / "batch_projects"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def default_batch_project_path() -> Path:
    return batch_projects_dir() / "new_batch_project.vbatch"


def save_batch_project_file(path_value, items: List[dict]) -> str:
    import json

    path = Path(path_value)
    data = {
        "format": "VADAFOK_BATCH_PROJECT",
        "version": 1,
        "items": items,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    try:
        temporary.write_text(
            json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return str(path)


def load_batch_project_file(path_value) -> List[dict]:
    import json

    path = Path(path_value)
    data = json.loads(path.read_text(encoding="utf-8"))

    if data.get("format") != "VADAFOK_BATCH_PROJECT":
        raise ValueError("Keine gültige VADAFOK Batch Project Datei.")

    items = data.get("items", [])
    if not isinstance(items, list):
        raise ValueError("Batch Project enthält keine gültige Item-Liste.")

    cleaned = []
    for item in items:
        if not isinstance(item, dict):
            continue
        values = item.get("values", {})
        cleaned.append({
            "template": str(item.get("template", "")),
            "output_name": str(item.get("output_name", "card")),
            "profile": str(item.get("profile", "Broadcast PNG")),
            "values": dict(values) if isinstance(values, dict) else {},
        })
    return cleaned
